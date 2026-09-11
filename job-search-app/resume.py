"""Conservative, local text extraction. No external model or profile storage."""
import io
import re
from matching import canonical

def read_document(name, data):
    if len(data)>5*1024*1024:
        raise ValueError('Maximum resume size is 5 MB.')
    if name.lower().endswith('.pdf'):
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        return [(f'Resume p. {i+1}',p.extract_text() or '') for i,p in enumerate(reader.pages)]
    if name.lower().endswith('.docx'):
        from docx import Document
        doc=Document(io.BytesIO(data))
        text='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
        return [('Resume DOCX',text)]
    if name.lower().endswith('.txt'):
        return [('Resume TXT',data.decode('utf-8-sig'))]
    raise ValueError('Upload PDF, DOCX or UTF-8 TXT.')

def evidence_for(term, documents):
    aliases={'SQL':['SQL','PostgreSQL','MySQL'],'Excel':['Excel','Microsoft Excel'],'Power BI':['Power BI','PowerBI'],'Git':['Git','GitHub']}
    for source,text in documents:
        for alias in aliases.get(term,[term]):
            flags=0 if alias=='R' else re.I
            m=re.search(r'(?<![\w&])'+re.escape(alias)+r'(?![\w&])',text,flags)
            if m:
                return {'Source':source,'Evidence':re.sub(r'\s+',' ',text[max(0,m.start()-65):m.end()+100])}
    return None

def extract(documents, vocabulary):
    evidence={s:evidence_for(s,documents) for s in vocabulary}
    evidence={s:e for s,e in evidence.items() if e}
    lines=[line.strip() for _,text in documents for line in text.splitlines() if line.strip()]
    education=next((x for x in lines if re.search(r'\b(university|college|UMKC|bachelor|master|Ph\.?D)\b',x,re.I)),'')
    location=next((m.group(0) for x in lines if (m:=re.search(r'\b[A-Z][a-zA-Z .-]+,\s*[A-Z]{2}\b',x))), '')
    return {'skills':', '.join(evidence),'education':education,'location':location,'evidence':evidence}
